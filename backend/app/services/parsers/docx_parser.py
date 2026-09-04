import io
import docx
from app.core.logging import get_logger

logger = get_logger(__name__)


def parse_docx(content: bytes) -> str:
    """Extracts plain text from DOCX bytes in-memory."""
    stream = io.BytesIO(content)
    try:
        document = docx.Document(stream)
        text_parts = []

        # Extract paragraphs
        for p in document.paragraphs:
            if p.text.strip():
                text_parts.append(p.text.strip())

        # Extract tables
        for table in document.tables:
            for row in table.rows:
                row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_texts:
                    text_parts.append(" | ".join(row_texts))

        return "\n\n".join(text_parts).strip()
    except Exception as e:
        logger.warning("docx_extraction_failed", error=str(e))
        return ""
